# -*- coding: utf-8 -*-
"""生成「禅道工时填报.xlsx」模板（数据区留空，示例只放在说明页，避免误提交）"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "templates", "禅道工时填报.xlsx")

HEAD = ["日期", "类型", "对象号", "工作内容", "消耗工时", "剩余工时", "测试包版本"]
WIDTH = [13, 10, 14, 46, 11, 11, 22]
HINT = ["留空=今天\n或 1.15", "task/bug", "如 123456\n或 任务#123456", "这条日志写什么", "小时，>0",
        "留空最好\n(自动计算)", "仅在 剩余=0\n时必填"]

hdr_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1D6E56")
hint_font = Font(name="微软雅黑", size=8, color="0F6E56")
hint_fill = PatternFill("solid", fgColor="E1F5EE")
body_font = Font(name="微软雅黑", size=10)
thin = Side(style="thin", color="B4B2A9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------------- Sheet1: 工时填报
ws = wb.active
ws.title = "工时填报"
ws.sheet_view.showGridLines = False

for i, (h, w, hint) in enumerate(zip(HEAD, WIDTH, HINT), start=1):
    c = ws.cell(1, i, h)
    c.font, c.fill, c.border = hdr_font, hdr_fill, border
    c.alignment = Alignment("center", "center")
    ws.column_dimensions[get_column_letter(i)].width = w
    hc = ws.cell(2, i, hint)
    hc.font, hc.fill, hc.border = hint_font, hint_fill, border
    hc.alignment = Alignment("center", "center", wrap_text=True)
ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 30

DATA_ROWS = 200
for r in range(3, 3 + DATA_ROWS):
    ws.row_dimensions[r].height = 18
    for ci in range(1, 8):
        cell = ws.cell(r, ci)
        cell.font, cell.border = body_font, border
        if ci == 1:
            cell.number_format = "@"     # 日期列设文本：否则 Excel 会把 "1.15" 存成数字 1.15
            cell.alignment = Alignment("center", "center")
        elif ci in (2, 3, 5, 6):
            cell.alignment = Alignment("center", "center")
        elif ci == 4:
            cell.alignment = Alignment("left", "center", wrap_text=True)

dv = DataValidation(type="list", formula1='"task,bug"', allow_blank=True,
                    showErrorMessage=True, errorTitle="类型不对",
                    error="只能是 task 或 bug")
ws.add_data_validation(dv)
dv.add(f"B3:B{3 + DATA_ROWS - 1}")

dv2 = DataValidation(type="decimal", operator="greaterThan", formula1="0",
                     allow_blank=True, showErrorMessage=True, errorTitle="消耗工时不对",
                     error="消耗工时必须大于 0")
ws.add_data_validation(dv2)
dv2.add(f"E3:E{3 + DATA_ROWS - 1}")

ws.freeze_panes = "A3"

# ---------------- Sheet2: 填写说明
ws2 = wb.create_sheet("填写说明")
ws2.sheet_view.showGridLines = False
ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 96


def line(r, k, v, bold=False, color="2C2C2A"):
    a = ws2.cell(r, 1, k)
    a.font = Font(name="微软雅黑", size=10, bold=True)
    a.alignment = Alignment("left", "top")
    b = ws2.cell(r, 2, v)
    b.font = Font(name="微软雅黑", size=10, bold=bold, color=color)
    b.alignment = Alignment("left", "top", wrap_text=True)
    return r + 1


r = 1
ws2.cell(r, 1, "禅道工时填报 · 填写说明").font = Font(name="微软雅黑", size=13, bold=True)
r += 2
r = line(r, "怎么用", "填完「工时填报」页 → 告诉 AI「帮我填禅道日志」，AI 会先跑 dry-run 给你看计划，确认后才真正写入。")
r += 1
r = line(r, "日期", "YYYY-MM-DD / 2026/1/15 / 1.15 / 今天 / 昨天。留空 = 今天。")
r = line(r, "类型", "task 或 bug（下拉选择）。")
r = line(r, "对象号", "任务/Bug 的编号。写成 123456、任务#123456、BUG#123456 都认。")
r = line(r, "工作内容", "这条日志写什么。中文随便写，会原样存进禅道的「工作」字段。")
r = line(r, "消耗工时", "本次花了多少小时，必须大于 0（禅道拒绝 0）。支持 1.5 这种小数。")
r += 1
r = line(r, "剩余工时", "留空【推荐】= 脚本按「任务当前剩余 − 本次消耗」自动算。\n"
                        "手填 = 你指定的绝对值，会覆盖任务当前剩余；同一任务多行按提交顺序累计扣减。\n"
                        "只有确实要把任务置为「完成」时才手填 0，见下方红线。")
r = line(r, "测试包版本", "只有当「剩余工时 = 0」时才必填，格式用等号：example-pkg=1.0.0-1")
r += 1
ws2.cell(r, 1, "红线").font = Font(name="微软雅黑", size=10, bold=True, color="A32D2D")
r += 1
for t in [
    "1. 写入前一定先看 dry-run 计划，确认后才写。",
    "2. 剩余工时填 0 → 任务立刻变「完成」，且必须给测试包版本。非真要交付，别填 0。",
    "3. 已提交过的行会记进台账（.ledger/），重复提交默认跳过，避免同一天刷出两条一样的工时。",
    "4. 每行 = 一条独立工时记录。同一天同一任务分几段做就写几行。",
    "5. 部分实例的工时接了外部工时系统同步：脚本可以新建工时，但改/删已有工时会被拦（<同步工时-…>），要改去网页。",
]:
    c = ws2.cell(r, 2, t)
    c.font = Font(name="微软雅黑", size=10, color="A32D2D")
    c.alignment = Alignment("left", "top", wrap_text=True)
    r += 1
r += 1
ws2.cell(r, 1, "示例").font = Font(name="微软雅黑", size=10, bold=True)
r += 1
ws2.cell(r, 1, "（示例，别抄进填报页）").font = Font(name="微软雅黑", size=9, color="888780")
r += 1
for row in [
    ("日期", "类型", "对象号", "工作内容", "消耗工时", "剩余工时", "测试包版本"),
    ("2026-01-15", "task", "123456", "处理模块 A 的功能开发与自测", "2", "", ""),
    ("2026-01-15", "task", "123457", "调研构建方案并输出结论", "3.5", "", ""),
    ("2026-01-15", "bug", "123458", "修复空指针问题并完成回归", "1.5", "0", "example-pkg=1.0.0-1"),
]:
    for i, v in enumerate(row, start=1):
        ws2.cell(r, i, v).font = Font(name="微软雅黑", size=9)
        ws2.column_dimensions[get_column_letter(i)].width = max(
            ws2.column_dimensions[get_column_letter(i)].width or 0, 16 if i != 4 else 30)
    r += 1

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved:", OUT)