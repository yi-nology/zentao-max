#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
禅道工时日志填报 CLI —— zentao-timelog skill 核心脚本

链路: 本机 python --(SSH -L 隧道)--> jumpbox(公司 VPN) --> zentao.example.com

子命令:
  tunnel            确保 SSH 转发隧道在线（不在则起）
  whoami            显示当前登录账号
  tasks [关键词]     列出"我的任务"（id/状态/已耗/剩余/名称）
  plan  <xlsx>      解析 Excel + 校验 + 计算 left，只输出计划，不写任何数据
  submit <xlsx>     执行填报（必须显式 --yes），带台账去重
  verify <日期>      回查某天已填报的工时

设计红线见 reference/chandao-api.md。
"""
import argparse
import atexit
import csv
import hashlib
import html as _html
import json
import os
import re
import socket
import ssl
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import http.cookiejar
from datetime import date as _date

# ---------------------------------------------------------------- 配置
# 全部可通过环境变量覆盖；默认值是占位符，使用前必须改成自己的环境。
HOST = os.environ.get("ZT_HOST", "zentao.example.com")   # 禅道域名（SNI / Host 用）
BACKEND = "127.0.0.1"                                     # 隧道走本机
PORT = int(os.environ.get("ZT_PORT") or "18443")          # 本机转发端口；直连时用 443
JUMPBOX = os.environ.get("ZT_JUMPBOX", "jumpbox")         # ~/.ssh/config 里的 alias；留空=直连
REALM = os.environ.get("ZT_REALM", "")                    # 统一认证域；留空则登录时不发送

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CRED_DEFAULT = os.path.expanduser(os.environ.get("ZT_CRED", "~/.zentao-credentials.env"))
CACHE_DIR = os.path.join(SKILL_DIR, ".cache")
LEDGER_DIR = os.path.join(SKILL_DIR, ".ledger")

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"


# ---------------------------------------------------------------- 小工具
def log(msg=""):
    print(msg, flush=True)


def cred_path():
    return os.environ.get("ZT_CRED", CRED_DEFAULT)


def load_cred(required=True):
    """读凭证文件（兼容 KEY=VALUE、去引号、跳过注释）"""
    path = cred_path()
    if not os.path.exists(path):
        if not required:
            return {}
        raise SystemExit(f"[x] 找不到凭证文件: {path}\n    可用 ZT_CRED 环境变量覆盖路径")
    out = {}
    with open(path, encoding="utf-8", errors="replace") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            out[k.strip()] = v.strip().strip('"').strip("'")
    return out


def apply_config():
    """把凭证文件里的 ZT_HOST / ZT_PORT / ZT_JUMPBOX / ZT_REALM 应用到模块常量。
    环境变量优先；文件里没有的键保留默认值。"""
    global HOST, PORT, JUMPBOX, REALM
    creds = load_cred(required=False)
    HOST = os.environ.get("ZT_HOST") or creds.get("ZT_HOST") or HOST
    if "ZT_JUMPBOX" in os.environ:
        JUMPBOX = os.environ["ZT_JUMPBOX"]
    else:
        JUMPBOX = creds.get("ZT_JUMPBOX", JUMPBOX)
    if "ZT_REALM" in os.environ:
        REALM = os.environ["ZT_REALM"]
    else:
        REALM = creds.get("ZT_REALM", REALM)
    port_raw = os.environ.get("ZT_PORT") or creds.get("ZT_PORT")
    if port_raw:
        try:
            PORT = int(port_raw)
        except ValueError:
            raise SystemExit(f"[x] ZT_PORT 不是数字: {port_raw!r}")
    else:
        PORT = 18443 if JUMPBOX else 443
    return creds


def port_open(port=None, host=None):
    port = PORT if port is None else port
    host = BACKEND if host is None else host
    with socket.socket() as s:
        s.settimeout(1.5)
        return s.connect_ex((host, port)) == 0


_TUNNEL_PROC = None          # 本进程拉起的隧道，随进程退出回收


def _tunnel_args():
    return ["ssh", "-N",
            "-o", "ExitOnForwardFailure=yes",
            "-o", "ServerAliveInterval=30",
            "-o", "ServerAliveCountMax=3",
            "-o", "ConnectTimeout=15",
            "-L", f"{PORT}:{HOST}:443",
            JUMPBOX]


def _cleanup_tunnel():
    global _TUNNEL_PROC
    if _TUNNEL_PROC is not None and _TUNNEL_PROC.poll() is None:
        try:
            _TUNNEL_PROC.terminate()
        except Exception:
            pass
    _TUNNEL_PROC = None


atexit.register(_cleanup_tunnel)


def ensure_tunnel(quiet=False):
    """
    保证 127.0.0.1:PORT 上有可用隧道。
      * 端口已在监听 -> 复用（可能是用户自己起的常驻隧道）
      * 否则本进程拉起 ssh -L，进程退出时自动回收

    注：agent 沙箱的 job object 会连带回收 ssh 子进程，跨命令无法常驻；
    所以每次调用都按需重建（约 1~2s），不要假设隧道一直在。
    """
    global _TUNNEL_PROC
    if not JUMPBOX:
        if not quiet:
            log("[=] ZT_JUMPBOX 为空，按直连模式访问")
        return
    if port_open():
        if not quiet:
            log(f"[=] 复用已在线隧道: {BACKEND}:{PORT} -> {HOST}:443")
        return
    if not quiet:
        log(f"[+] 拉起隧道: {BACKEND}:{PORT} -> {HOST}:443 (via {JUMPBOX})")
    _TUNNEL_PROC = subprocess.Popen(_tunnel_args(), stdout=subprocess.DEVNULL,
                                    stderr=subprocess.DEVNULL, stdin=subprocess.DEVNULL)
    for _ in range(40):
        time.sleep(0.5)
        if port_open():
            if not quiet:
                log("[=] 隧道就绪")
            return
        if _TUNNEL_PROC.poll() is not None:
            break
    _cleanup_tunnel()
    raise SystemExit("[x] 隧道 20s 内未就绪。多半是跳板机的 VPN 掉线 —— "
                     "让用户重连 VPN 后重试。")


# ---------------------------------------------------------------- HTTP 会话
class Zentao:
    """
    通过隧道访问禅道的极简客户端。
    要点:
      * 关闭所有代理（本机 HTTP_PROXY 会抢走 CONNECT，必须显式绕开）
      * 走隧道时把 HOST 的 DNS 解析劫持到 127.0.0.1（保证 SNI=HOST 走隧道）
      * 表单以 UTF-8 urlencode 提交（不可用 EscapeDataString/%XX 二次编码）
    """

    def __init__(self):
        if JUMPBOX:
            self.base = f"https://{HOST}:{PORT}"
        else:
            self.base = f"https://{HOST}" if PORT in (443, 0) else f"https://{HOST}:{PORT}"
        self.uid = None
        self.account = None
        self.realname = None
        if JUMPBOX:
            self._patch_dns()
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
        self.jar = http.cookiejar.CookieJar()
        self.op = urllib.request.build_opener(
            urllib.request.ProxyHandler({}),
            urllib.request.HTTPSHandler(context=ctx),
            urllib.request.HTTPCookieProcessor(self.jar),
        )

    @staticmethod
    def _patch_dns():
        if getattr(socket.getaddrinfo, "_zt_patched", False):
            return
        orig = socket.getaddrinfo

        def patched(host, *a, **kw):
            return orig(BACKEND if host == HOST else host, *a, **kw)

        patched._zt_patched = True
        socket.getaddrinfo = patched

    def _open(self, path, data=None, timeout=60):
        url = self.base + path
        headers = {"User-Agent": USER_AGENT}
        body = None
        if data is not None:
            body = urllib.parse.urlencode(data, encoding="utf-8").encode("utf-8")
            headers["Content-Type"] = "application/x-www-form-urlencoded; charset=utf-8"
            headers["X-Requested-With"] = "XMLHttpRequest"
        req = urllib.request.Request(url, data=body, headers=headers)
        try:
            with self.op.open(req, timeout=timeout) as r:
                return r.status, r.read().decode("utf-8", "replace")
        except urllib.error.HTTPError as e:
            return e.code, e.read().decode("utf-8", "replace")
        except Exception as e:
            return 0, f"CONNERR:{e}"

    def get(self, path, timeout=60):
        return self._open(path, None, timeout)

    def post(self, path, data, timeout=90):
        return self._open(path, data, timeout)

    def login(self):
        creds = load_cred(required=False)
        user = os.environ.get("ZT_USERNAME") or creds.get("ZT_USERNAME", "")
        pwd = os.environ.get("ZT_PASSWORD") or creds.get("ZT_PASSWORD", "")
        if not user or not pwd:
            raise SystemExit(f"[x] 凭证文件缺 ZT_USERNAME / ZT_PASSWORD ({cred_path()})")
        st, _ = self.get("/user-login-Lw==.html")
        if st != 200:
            raise SystemExit(f"[x] 拿不到登录页 (HTTP {st})，先跑 `zt.py tunnel` 并确认 VPN/网络在线")
        st, rnd = self.get("/user-refreshRandom.html")
        form = {
            "account": user, "password": pwd, "verifyRand": rnd.strip(),
            "keepLogin": "0", "referer": "/",
        }
        if REALM:
            form["realm"] = REALM
        st, resp = self.post("/user-login.json", form)
        if '"success"' not in resp:
            raise SystemExit(f"[x] 登录失败: {resp[:300]}")
        m = re.search(r'"id":(\d+)', resp)
        self.uid = int(m.group(1)) if m else None
        self.account = user
        m = re.search(r'"realname":"([^"]*)"', resp)
        self.realname = m.group(1).encode().decode("unicode_escape") if m else ""
        return self

    def my_tasks(self):
        st, txt = self.get("/my-work.json")
        if st != 200 or '"success"' not in txt:
            return []
        try:
            inner = json.loads(json.loads(txt)["data"])
            tasks = inner.get("tasks", [])
            for t in tasks:                      # 名称里的 &amp; 等实体还原
                if isinstance(t.get("name"), str):
                    t["name"] = _html.unescape(t["name"])
            return tasks
        except Exception:
            return []

    def task_by_id(self, tid, tasks=None):
        for t in (tasks if tasks is not None else self.my_tasks()):
            if int(t.get("id", 0)) == int(tid):
                return t
        return None


# ---------------------------------------------------------------- 对象号解析
def parse_object(raw):
    """
    从用户写法里提取 (类型, 数字ID)。支持:
      123456 / 任务#123456 / task#123456 / task_123456 / BUG#123456 / bug123456
    返回 (type, id) ；无法解析抛 ValueError
    """
    if raw is None:
        raise ValueError("对象号为空")
    s = str(raw).strip()
    if not s:
        raise ValueError("对象号为空")
    low = s.lower()
    if "bug" in low or "缺陷" in s:
        otype = "bug"
    elif "task" in low or "任务" in s:
        otype = "task"
    else:
        otype = "task"
    m = re.search(r"(\d+)", s)
    if not m:
        raise ValueError(f"对象号里找不到数字: {s!r}（应形如 123456 或 任务#123456）")
    return otype, int(m.group(1))


def _mk_date(y, mo, d, raw):
    try:
        return _date(y, mo, d).isoformat()
    except ValueError:
        raise ValueError(f"日期不存在: {raw!r}（{y}-{mo}-{d} 不是合法日期）")


def norm_date(raw):
    """
    支持多种写法（写错会明确报错，不要静默猜）：
      2026-09-18 / 2026/9/18 / 2026.9.18          完整日期
      09-18 / 9/18 / 9.18 / 9.16                  省略年份 -> 当年
      今天 / 今日 / today / 昨天 / yesterday
      留空 = 今天
      Excel 把 "9.16" 自动转成的数字 9.16          按 月.日 解释
    """
    import datetime as _dt
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return _date.today().isoformat()
    if isinstance(raw, _dt.datetime):
        return raw.date().isoformat()
    if isinstance(raw, _date):
        return raw.isoformat()
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        f = float(raw)
        # Excel 日期序列号（1900 起）
        if 20000 < f < 80000:
            return (_dt.date(1899, 12, 30) + _dt.timedelta(days=int(f))).isoformat()
        s = ("%.4f" % f).rstrip("0").rstrip(".")       # 9.16 -> "9.16"
        m = re.match(r"^(\d{1,2})\.(\d{1,2})$", s)
        if m:
            return _mk_date(_date.today().year, int(m.group(1)), int(m.group(2)), raw)
        raise ValueError(f"日期识别不了: {raw!r}（数字形式只支持 月.日，如 9.18）")

    s = str(raw).strip()
    low = s.lower()
    if low in ("今天", "今日", "today", "now"):
        return _date.today().isoformat()
    if low in ("昨天", "昨日", "yesterday") or "昨天" in s or "昨日" in s:
        return (_date.today() - _dt.timedelta(days=1)).isoformat()
    s = s.split(" ")[0].split("T")[0].strip()
    s = re.sub(r"年|月", "-", s).replace("日", "").replace("/", "-").replace(".", "-")
    s = re.sub(r"-+", "-", s).strip("-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if m:
        return _mk_date(int(m.group(1)), int(m.group(2)), int(m.group(3)), raw)
    m = re.match(r"^(\d{1,2})-(\d{1,2})$", s)          # 省略年份
    if m:
        return _mk_date(_date.today().year, int(m.group(1)), int(m.group(2)), raw)
    raise ValueError(f"日期格式不对: {raw!r}（应为 YYYY-MM-DD / 9.18，留空=今天）")


def to_num(raw, field, row):
    if raw is None or str(raw).strip() == "":
        return None
    try:
        return float(str(raw).strip())
    except ValueError:
        raise ValueError(f"第 {row} 行「{field}」不是数字: {raw!r}")


# ---------------------------------------------------------------- Excel 解析
HEADER_ALIASES = {
    "date":   ["日期", "date", "工作日期"],
    "type":   ["类型", "对象类型", "type", "objecttype"],
    "oid":    ["对象号", "任务号", "编号", "id", "对象id", "任务/bug号"],
    "work":   ["工作内容", "内容", "工作描述", "work", "日志", "日志内容"],
    "consumed": ["消耗工时", "消耗(h)", "工时", "消耗", "consumed", "小时"],
    "left":   ["剩余工时", "剩余(h)", "剩余", "left", "剩余小时"],
    "pkg":    ["测试包版本", "包版本", "testpackageversion", "包名=版本号"],
}


def _norm_header(s):
    return re.sub(r"[\s（）()\[\]【】:：]", "", str(s or "")).lower()


LEGEND_MARKERS = ("留空=今天", "task/bug", "或任务#", "这条日志写什么",
                  "小时，>0", "留空=自动", "仅在剩余", "包名=版本")


def _is_legend(vals):
    """识别模板自带的说明行（「留空=今天 / task/bug / 这条日志写什么」），命 2 个以上即整行跳过"""
    hits = 0
    for v in vals.values():
        s = re.sub(r"\s+", "", str(v)) if v is not None else ""
        if not s:
            continue
        if any(re.sub(r"\s+", "", mk) in s for mk in LEGEND_MARKERS):
            hits += 1
    return hits >= 2


def read_excel(path):
    """
    读 Excel 第一个工作表 -> rows(list[dict])
    期望表头（顺序随意、可少 剩余工时/测试包版本，其余必需）:
      日期 | 类型 | 对象号 | 工作内容 | 消耗工时 | 剩余工时 | 测试包版本
    """
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("[x] 缺 openpyxl。先执行:  pip install openpyxl")
    if not os.path.exists(path):
        raise SystemExit(f"[x] 找不到 Excel: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    # 定位表头行（前 10 行内找命中最多的一行）
    hdr_row, colmap, best = None, {}, 0
    for r in range(1, min(11, ws.max_row + 1)):
        found = {}
        for c in range(1, ws.max_column + 1):
            h = _norm_header(ws.cell(r, c).value)
            if not h:
                continue
            for key, names in HEADER_ALIASES.items():
                if key in found:
                    continue
                if any(_norm_header(n) == h or _norm_header(n) in h for n in names):
                    found[key] = c
                    break
        if len(found) > best:
            best, hdr_row, colmap = len(found), r, found
    if not colmap or not {"oid", "work", "consumed"} <= set(colmap):
        raise SystemExit("[x] 表头识别失败。第一张表需含列：日期/类型/对象号/工作内容/消耗工时 "
                         "（可选 剩余工时、测试包版本）。参考 templates/禅道工时填报.xlsx")

    rows = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        def cell(key):
            c = colmap.get(key)
            return ws.cell(r, c).value if c else None

        vals = {k: cell(k) for k in colmap}
        if all(v in (None, "") for v in vals.values()):
            continue
        raw_work = cell("work")
        if isinstance(raw_work, str) and raw_work.strip().startswith("#"):
            continue  # 注释行
        if _is_legend(vals):
            continue  # 模板自带的说明行（第 2 行那种）
        if cell("oid") in (None, "") and cell("work") in (None, ""):
            continue
        rows.append({"excel_row": r, "date": cell("date"), "type": cell("type"),
                     "oid": cell("oid"), "work": raw_work,
                     "consumed": cell("consumed"), "left": cell("left"), "pkg": cell("pkg")})
    if not rows:
        raise SystemExit("[x] Excel 里没有有效数据行")
    return rows


# ---------------------------------------------------------------- 台账（防重复填报）
def ledger_path(d):
    return os.path.join(LEDGER_DIR, f"{d[:7]}.csv")


def ledger_key(d, otype, oid, work):
    h = hashlib.sha1((work or "").strip().encode("utf-8")).hexdigest()[:10]
    return f"{d}|{otype}|{oid}|{h}"


def ledger_load(d):
    seen = {}
    os.makedirs(LEDGER_DIR, exist_ok=True)
    p = ledger_path(d)
    if not os.path.exists(p):
        return seen
    with open(p, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            seen[row.get("key", "")] = row
    return seen


def ledger_load_all():
    """把 .ledger 下所有月份的台账并成一份。跨月填报时只读单月会漏判重复。"""
    seen = {}
    os.makedirs(LEDGER_DIR, exist_ok=True)
    for fn in sorted(os.listdir(LEDGER_DIR)):
        if not fn.endswith(".csv"):
            continue
        try:
            with open(os.path.join(LEDGER_DIR, fn), encoding="utf-8-sig", newline="") as f:
                for row in csv.DictReader(f):
                    k = row.get("key", "")
                    if k:
                        seen[k] = row
        except Exception:
            continue
    return seen


def ledger_append(d, rec):
    p = ledger_path(d)
    new = not os.path.exists(p)
    os.makedirs(LEDGER_DIR, exist_ok=True)
    with open(p, "a", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["key", "at", "date", "type", "oid", "consumed",
                                          "left", "work", "resp", "effort_id"])
        if new:
            w.writeheader()
        w.writerow(rec)
        f.flush()
        os.fsync(f.fileno())        # 台账是防重复填报的唯一依据，必须落盘


# ---------------------------------------------------------------- 计划构建
def build_plan(zt, rows, force=False):
    """
    把 Excel 行变成可执行计划，并做全部校验。
    返回 (plan, errors, warnings)。plan 每项含 date/type/oid/work/consumed/left/pkg/danger
    """
    tasks = zt.my_tasks()
    by_id = {}
    for t in tasks:
        try:
            by_id[int(t.get("id", 0))] = t
        except Exception:
            pass

    plan, errors, warns = [], [], []
    running = {}          # (otype, oid) -> 提交完前面各行之后该对象的剩余工时
    seen = ledger_load_all()   # 台账（按月存，跨月也要读全）
    for r in rows:
        n = r["excel_row"]
        try:
            d = norm_date(r["date"])
        except ValueError as e:
            errors.append(f"第 {n} 行: {e}")
            continue
        try:
            otype, oid = parse_object(r["oid"])
        except ValueError as e:
            errors.append(f"第 {n} 行: {e}")
            continue
        # 显式类型列可覆盖
        tcol = str(r.get("type") or "").strip().lower()
        if tcol:
            if "bug" in tcol or "缺陷" in tcol:
                otype = "bug"
            elif "task" in tcol or "任务" in tcol:
                otype = "task"
            else:
                errors.append(f"第 {n} 行: 类型只支持 task / bug，收到 {r['type']!r}")
                continue

        work = str(r.get("work") or "").strip()
        if not work:
            errors.append(f"第 {n} 行: 工作内容为空")
            continue
        try:
            consumed = to_num(r.get("consumed"), "消耗工时", n)
        except ValueError as e:
            errors.append(str(e))
            continue
        if consumed is None or consumed <= 0:
            errors.append(f"第 {n} 行: 消耗工时必须 > 0（禅道拒绝 0）")
            continue
        if consumed > 24:
            warns.append(f"第 {n} 行: 单日消耗 {consumed}h 超过 24h，确认这是你要的吗？")

        try:
            left_in = to_num(r.get("left"), "剩余工时", n)
        except ValueError as e:
            errors.append(str(e))
            continue
        pkg = str(r.get("pkg") or "").strip() or None

        cur_left = cur_consumed = None
        task = by_id.get(oid) if otype == "task" else None
        if task is not None:
            try:
                cur_left = float(task.get("left") or 0)
                cur_consumed = float(task.get("consumed") or 0)
            except Exception:
                cur_left = cur_consumed = None

        # 同一个对象出现多行时，剩余工时是【累计递减】的：禅道按「旧剩余 − 本次消耗」
        # 逐条往下扣。dry-run 必须用同一套口径，否则第 2 行以后基准值是错的（会漏判负数）。
        okey = (otype, oid)
        if okey not in running and cur_left is not None:
            running[okey] = cur_left
        run_left = running.get(okey)

        # 台账命中的行不会真正提交 -> 不能让它推进累计基准（否则后面几行的基准会多扣一次）
        key = ledger_key(d, otype, oid, work)
        dup = seen.get(key) if not force else None

        danger = False   # danger=True 表示会把对象置为 done（left=0）
        if left_in is None:
            # 未填剩余工时 -> 自然语义：当前剩余 - 本次消耗，但不允许自动置 0
            if run_left is None:
                errors.append(f"第 {n} 行: 对象 {oid} 不在「我的任务」里，无法自动推算剩余工时，"
                              f"请在「剩余工时」列显式填写")
                continue
            calc = round(run_left - consumed, 2)
            if calc <= 0:
                errors.append(
                    f"第 {n} 行: 累计剩余 {run_left}h - 本次 {consumed}h = {calc}h ≤ 0，"
                    f"会把任务置为 done。若确实要完成，请在「剩余工时」列显式写 0 并在"
                    f"「测试包版本」列填「包名=版本号」")
                continue
            left = calc
        else:
            left = left_in
            if left == 0:
                danger = True
                if not pkg:
                    errors.append(f"第 {n} 行: 剩余工时=0 会把任务置为 done，"
                                  f"此时「测试包版本」必填（格式 包名=版本号，如 example-pkg=1.0.0-1）")
                    continue
                if "=" not in pkg:
                    errors.append(f"第 {n} 行: 测试包版本要用等号，如 example-pkg=1.0.0-1（收到 {pkg!r}）")
                    continue

        if not dup:
            running[okey] = left      # 只有会真正提交的行才推进基准

        plan.append({"excel_row": n, "date": d, "type": otype, "oid": oid, "work": work,
                     "consumed": consumed, "left": left,
                     "left_explicit": left_in is not None,    # 仅用于展示：是你指定的还是自动推的
                     "pkg": pkg, "danger": danger, "cur_left": run_left,
                     "cur_consumed": cur_consumed, "key": key, "dup": dup,
                     "execution": (task or {}).get("execution"),
                     "task_name": (task or {}).get("name", "")})

    if any(p.get("dup") for p in plan):
        warns.append("有 %d 行在台账里已存在（疑似重复填报），默认会跳过；要强制重填加 --force"
                     % sum(1 for p in plan if p.get("dup")))

    # 同一对象多行 -> 明确告诉用户剩余工时是累计扣减的，避免误读成"每条都从初值扣"
    per_obj = {}
    for p in plan:
        if p.get("dup"):
            continue
        per_obj.setdefault((p["type"], p["oid"]), []).append(p)
    for (ot, oid), ps in sorted(per_obj.items()):
        if len(ps) > 1 and ps[0]["cur_left"] is not None:
            warns.append("对象 %s-%s 共 %d 行，剩余工时按提交顺序累计扣减：%g → %g"
                         % (ot, oid, len(ps), ps[0]["cur_left"], ps[-1]["left"]))
    return plan, errors, warns


# ---------------------------------------------------------------- 渲染
def print_plan(plan, errors, warns, title="填报计划"):
    log("")
    log("=" * 78)
    log(f"  {title}")
    log("=" * 78)
    tot = 0.0
    for p in plan:
        tot += p["consumed"]
        tag = "  [置为DONE]" if p["danger"] else ""
        dup = "  [台账已有→跳过]" if p.get("dup") else ""
        left_txt = (f"剩余→{p['left']:>6.2f}h（你指定）" if p["left_explicit"]
                    else f"剩余→{p['left']:>6.2f}h（自动:{p['cur_left']:g}−{p['consumed']:g}）")
        log(f"  第{p['excel_row']:>3}行  {p['date']}  {p['type']:<4} #{p['oid']:<8} "
            f"{p['consumed']:>5.2f}h  {left_txt}{tag}{dup}")
        log(f"           内容: {p['work']}")
        if p["pkg"]:
            log(f"           测试包版本: {p['pkg']}")
        if p["task_name"]:
            log(f"           任务: {p['task_name']}")
    log("-" * 78)
    log(f"  共 {len(plan)} 条，合计 {tot:.2f} h")
    if warns:
        log("\n  [警告]")
        for w in warns:
            log(f"    - {w}")
    if errors:
        log("\n  [错误——必须先修掉]")
        for e in errors:
            log(f"    ! {e}")
    log("=" * 78)


# ---------------------------------------------------------------- 提交
def submit_one(zt, p):
    """
    提交一条工时。要点（2026-09-18 两次真实写入实测后固化）：

      * **`left[1]` 必须发，而且发的是绝对剩余值。** 本条实例校验「日期N : 剩余不能为空」，
        整条表单会被直接拒绝（HTTP 仍 200）。语义是把任务的「预计剩余」**赋**成这个值，
        实测发 25 就落 25（consumed 同步 +5）。
        取值 = 提交前该对象的剩余 − 本次消耗，同一对象多行时逐行累计（见 build_plan.running）。
      * `execution[1]` 一并带上：本实例的工时同步插件会校验「对象为任务时执行不能为空」。
      * 成功的响应是 `<script>if(parent !== window) parent.location.reload(true);</script>`；
        失败时 HTTP 仍是 200，但 body 里是 `{"status":"success","data":"{\\"message\\":\\"…\\"}"}`
        —— **AJAX 外壳的 status 恒为 success，只有 data 里的 message 才是业务结果**，
        所以判定必须穿透外壳（_judge 负责）。
    """
    path = f"/effort-createForObject-{p['type']}-{p['oid']}.json"
    data = {
        "id[1]": "1",
        "dates[1]": p["date"],
        "objectType[1]": p["type"],
        "objectID[1]": str(p["oid"]),
        "work[1]": p["work"],
        "consumed[1]": ("%g" % p["consumed"]),
        "left[1]": ("%g" % p["left"]),
    }
    if p.get("execution"):
        data["execution[1]"] = str(p["execution"])
    if p.get("pkg"):
        data["testPackageVersion"] = p["pkg"]

    st, resp = zt.post(path, data)
    ok, msg = _judge(resp)
    return ok, st, resp, msg


def _judge(resp):
    """判定禅道响应，返回 (ok, 人话消息)。
    形态有三种，实测都遇到过：
      1) 成功: <html>…<script>if(parent !== window) parent.location.reload(true); </script>
      2) 表单校验失败: {"status":"success","data":"{\\"message\\":\\"日期1 : 剩余不能为空。\\"}"}
         —— 注意最外层 status 是 "success"，只看这个字符串会把失败判成成功。
      3) 业务失败: {"result":"fail","message":"…"} / alert('…')
    """
    low = re.sub(r"\s+", " ", resp or "")

    # (2) AJAX 外壳：把 data 这层 JSON 字符串解出来看 message
    m = re.search(r'"data"\s*:\s*"((?:[^"\\]|\\.)*)"', low)
    if m:
        raw = m.group(1)
        try:
            inner = json.loads('"' + raw + '"')
        except Exception:
            inner = raw
        try:
            obj = json.loads(inner)
        except Exception:
            obj = None
        if isinstance(obj, dict):
            if obj.get("message"):
                return False, str(obj["message"])[:200]
            if obj.get("result") == "fail":
                return False, str(obj.get("message") or "fail")[:200]
        elif inner.strip() and "<script" not in inner:
            return False, inner.strip()[:200]

    # (3) 显式失败
    for pat in (r'"result"\s*:\s*"fail"', r'"result"\s*:\s*"error"', r'"status"\s*:\s*"fail"'):
        if re.search(pat, low):
            mm = re.search(r'"message"\s*:\s*"([^"]{0,200})"', low)
            return False, (mm.group(1) if mm else low[:200])
    m = re.search(r"alert\(\s*'([^']{0,200})'", low)
    if m:
        return False, m.group(1)

    # (1) 成功
    if ("parent.location.reload" in low) or ('"result":"success"' in low.replace(" ", "")):
        return True, "ok"
    return False, (low[:200] or "空响应")


def fetch_efforts(zt):
    """拉「地盘-我的日志」全部工时（本实例该接口是分页 JSON，data 里是转义过的字符串）"""
    st, txt = zt.get("/my-effort.json")
    if st != 200:
        return None
    try:
        outer = json.loads(txt)
        inner = outer.get("data")
        if isinstance(inner, str):
            inner = json.loads(inner)
        return (inner or {}).get("efforts", [])
    except Exception:
        return None


def verify_day(zt, d):
    """回查某天的工时。返回 (记录列表, 覆盖提示或 None)。

    ⚠️ `/my-effort.json` 无视所有查询参数（date/pageID/account 都试过），**恒返回最新 20 条**。
    所以日期早于这 20 条的下界时，结果是不可信的 —— 返回提示而不是让调用方误报「0 条」。
    """
    eff = fetch_efforts(zt)
    if eff is None:
        return None, "回查接口无响应"
    dates = sorted({str(e.get("date", ""))[:10] for e in eff if e.get("date")})
    mine = [e for e in eff if str(e.get("date", ""))[:10] == d]
    tip = None
    if dates and d < dates[0]:
        tip = (f"/my-effort.json 只返回最新 {len(eff)} 条（最早 {dates[0]}），"
               f"查不到 {d} 不代表没填 —— 要确认请看 /task-view-{{任务ID}}.html 的变更历史")
    return mine, tip


def cmd_tasks(args):
    ensure_tunnel(quiet=True)
    zt = Zentao().login()
    log(f"[=] 已登录: {zt.realname} ({zt.account}) uid={zt.uid}")
    tasks = zt.my_tasks()
    kw = (args.keyword or "").strip().lower()
    rows = [t for t in tasks if not kw or kw in str(t.get("name", "")).lower()
            or kw in str(t.get("id", ""))]
    log(f"\n{'ID':>9}  {'状态':<8} {'已耗':>6} {'剩余':>6}  {'执行':>7}  名称")
    log("-" * 100)
    for t in rows:
        log(f"{t.get('id'):>9}  {str(t.get('status')):<8} {float(t.get('consumed') or 0):>6.1f} "
            f"{float(t.get('left') or 0):>6.1f}  {t.get('execution'):>7}  {t.get('name')}")
    log(f"\n共 {len(rows)} 个任务（{len(tasks)} 个在「我的任务」里）")


def cmd_plan(args):
    ensure_tunnel(quiet=True)
    zt = Zentao().login()
    log(f"[=] 已登录: {zt.realname} ({zt.account}) uid={zt.uid}")
    rows = read_excel(args.xlsx)
    plan, errors, warns = build_plan(zt, rows, force=args.force)
    print_plan(plan, errors, warns, title=f"填报计划（dry-run，未写入）: {os.path.basename(args.xlsx)}")
    if errors:
        raise SystemExit(2)
    log("\n[i] 确认无误后执行:  zt.py submit \"%s\" --yes" % args.xlsx)


def cmd_submit(args):
    ensure_tunnel(quiet=True)
    zt = Zentao().login()
    log(f"[=] 已登录: {zt.realname} ({zt.account}) uid={zt.uid}")
    rows = read_excel(args.xlsx)
    plan, errors, warns = build_plan(zt, rows, force=args.force)
    print_plan(plan, errors, warns, title="即将填报（真实写入禅道）")
    if errors:
        raise SystemExit("[x] 有校验错误，已中止，未写入任何数据")
    todo = [p for p in plan if not p.get("dup")]
    if not todo:
        raise SystemExit("[=] 全部行都在台账里，无需填报")
    if not args.yes:
        log(f"\n[i] 未加 --yes，本次不写入。确认后重跑:  zt.py submit \"{args.xlsx}\" --yes")
        return

    log(f"\n[>] 开始填报 {len(todo)} 条 ...")
    ok_n = fail_n = 0
    written = []
    for p in todo:
        ok, st, resp, msg = submit_one(zt, p)
        flag = "OK  " if ok else "FAIL"
        log(f"  [{flag}] {p['date']} {p['type']} #{p['oid']} {p['consumed']}h  "
            f"剩余→{p['left']:g}h  HTTP {st}" + ("" if ok else f"  <- {msg}"))
        if not ok:
            fail_n += 1
            continue
        ok_n += 1
        written.append(p)
        ledger_append(p["date"], {
            "key": p["key"], "at": time.strftime("%Y-%m-%d %H:%M:%S"), "date": p["date"],
            "type": p["type"], "oid": p["oid"], "consumed": "%g" % p["consumed"],
            "left": "%g" % p["left"], "work": p["work"],
            "resp": re.sub(r"\s+", " ", resp or "")[:160],
            "effort_id": "",
        })
        time.sleep(0.4)

    log(f"\n[=] 完成: 成功 {ok_n} 条，失败 {fail_n} 条（台账已记 {LEDGER_DIR}）")
    if fail_n:
        log("[!] 有失败行：它们【没有】记进台账，修掉原因后可安全重跑（不会重复填报）")

    # 台账落盘校验：台账是唯一防重复依据，没写进去下次就会重复填报
    if written:
        miss = []
        for d in sorted({p["date"] for p in written}):
            seen = ledger_load(d)
            for p in written:
                if p["date"] == d and p["key"] not in seen:
                    miss.append(f"{d} {p['type']}#{p['oid']}  {str(p['work'])[:24]}")
        if miss:
            log("\n[!!] 台账未落盘 —— 去重已失效，切勿重跑 submit（会重复填报）:")
            for m in miss:
                log(f"     - {m}")
            log(f"     请检查写权限: {LEDGER_DIR}")

    # ---- 回读校验：以禅道侧真实数值为准（成功响应不代表数值落对）----
    if written:
        log("\n[>] 回读校验（禅道侧真实状态）:")
        tasks = zt.my_tasks()
        tmap = {}
        for t in tasks:
            try:
                tmap[int(t.get("id"))] = t
            except Exception:
                pass
        base = {}       # oid -> (消耗前, 剩余前)
        delta = {}      # oid -> 本次总消耗
        last_left = {}  # oid -> 期望剩余（最后一条）
        for p in written:
            oid = int(p["oid"])
            if p["type"] != "task":
                base.setdefault(oid, (None, None))
            elif oid not in base:
                base[oid] = (p["cur_consumed"], p["cur_left"])
            delta[oid] = delta.get(oid, 0) + p["consumed"]
            last_left[oid] = p["left"]
        for oid in sorted({int(p["oid"]) for p in written if p["type"] == "task"}):
            t = tmap.get(oid)
            if t is None:
                log(f"  [?] task #{oid}: 不在「我的任务」里，无法回读，请网页确认")
                continue
            cons = float(t.get("consumed") or 0)
            left = float(t.get("left") or 0)
            b_cons, b_left = base.get(oid, (None, None))
            exp_cons = (b_cons if b_cons is not None else cons - delta[oid]) + delta[oid]
            exp_left = last_left.get(oid, left)
            okc = abs(cons - exp_cons) < 0.01
            okl = abs(left - exp_left) < 0.01
            log(f"  [{'OK' if okc else '!!'}] task #{oid}  已耗 {cons:g}h（期望 {exp_cons:g}h）"
                f"（写前 {b_cons:g}）   剩余 {left:g}h（期望 {exp_left:g}h）（写前 {b_left:g}）")
            if not okc:
                log(f"        ! 已耗没落对！请到禅道网页核对 task {oid} 的工时记录，"
                    f"必要时删除重填（.ledger 已记录本次，重跑会被跳过）")
            elif not okl:
                log(f"        ! 工时正确，但「预计剩余」禅道落的是 {left:g}h 而非 {exp_left:g}h"
                    f"（禅道不保证采纳外部提交的 left）。想改到 {exp_left:g}h 请到网页 "
                    f"task-edit-{oid}.html 手工填")

    days = sorted({p["date"] for p in todo})
    log("\n[>] 回查当日工时记录:")
    for d in days:
        mine, tip = verify_day(zt, d)
        if mine is None:
            log(f"  {d}: 回查失败（不影响已提交结果，可稍后单独 verify）")
            continue
        log(f"  {d}: 禅道当日共 {len(mine)} 条工时记录")
        for e in mine[:20]:
            log(f"      #{e.get('id')} {e.get('objectType')}-{e.get('objectID')} "
                f"{e.get('consumed')}h  {str(e.get('work',''))[:40]}")
        if tip:
            log(f"      [!] {tip}")


def cmd_verify(args):
    ensure_tunnel(quiet=True)
    zt = Zentao().login()
    d = norm_date(args.date)
    mine, tip = verify_day(zt, d)
    if mine is None:
        raise SystemExit("[x] 回查失败：接口无响应")
    log(f"[=] {d} 禅道工时记录 {len(mine)} 条：")
    tot = 0.0
    for e in mine:
        try:
            tot += float(e.get("consumed") or 0)
        except Exception:
            pass
        log(f"  #{e.get('id')}  {e.get('objectType')}-{e.get('objectID')}  "
            f"{e.get('consumed')}h  {str(e.get('work',''))[:50]}")
    log(f"  合计 {tot:.2f} h")
    if tip:
        log(f"[!] {tip}")


def cmd_whoami(args):
    ensure_tunnel(quiet=True)
    zt = Zentao().login()
    log(f"[=] {zt.realname} ({zt.account}) uid={zt.uid}  @ {HOST}")


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(prog="zt.py", description="禅道工时日志填报")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sub.add_parser("tunnel", help="确保 SSH 隧道在线").set_defaults(func=lambda a: ensure_tunnel())

    sub.add_parser("whoami", help="显示当前登录账号").set_defaults(func=cmd_whoami)

    p = sub.add_parser("tasks", help="列出我的任务")
    p.add_argument("keyword", nargs="?", help="按名称/编号过滤")
    p.set_defaults(func=cmd_tasks)

    p = sub.add_parser("plan", help="解析 Excel 并输出计划（不写入）")
    p.add_argument("xlsx")
    p.add_argument("--force", action="store_true", help="忽略台账去重")
    p.set_defaults(func=cmd_plan)

    p = sub.add_parser("submit", help="执行填报")
    p.add_argument("xlsx")
    p.add_argument("--yes", action="store_true", help="确认真实写入")
    p.add_argument("--force", action="store_true", help="忽略台账去重")
    p.set_defaults(func=cmd_submit)

    p = sub.add_parser("verify", help="回查某日已填工时")
    p.add_argument("date", nargs="?", help="YYYY-MM-DD，留空=今天")
    p.set_defaults(func=cmd_verify)

    args = ap.parse_args()
    apply_config()
    args.func(args)


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except KeyboardInterrupt:
        log("\n[!] 已中断")
        sys.exit(130)
